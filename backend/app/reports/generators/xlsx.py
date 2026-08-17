"""XLSX writer (FR-11): the same defect level rows as `app.reports.generators.csv`, written
with `openpyxl` and with the header row frozen and filterable, since the point of asking for
the spreadsheet rather than the CSV is to sort and pivot in place.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.reports.document import ReportDocument
from app.reports.generators.csv import header, rows

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# Wide enough for the identity and narrative columns, default elsewhere; index-keyed so the
# widths stay attached to `csv._COLUMN_KEYS`'s order rather than to a duplicated name list.
_COLUMN_WIDTHS = {1: 18, 2: 16, 3: 38, 8: 16, 10: 34, 14: 60, 15: 44, 16: 44}


def write(document: ReportDocument, path: Path) -> int:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inspections"

    columns = header(document.language)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    written = 0
    for row in rows(document):
        sheet.append(row)
        written += 1

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(columns)).coordinate}"
    for index, width in _COLUMN_WIDTHS.items():
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    workbook.save(path)
    return written
